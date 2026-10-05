#!/usr/bin/env python3
"""Archive and inspect NKR attachment files for concession coordinate registers.

Inputs:
  map/data/nkr-party-documents-v1.json
  map/data/pending-concessions-v1.json
  map/data/concession-boundaries-official-source-v2.json

Outputs:
  map/data/nkr-attachment-index-v1.json
  map/data/coordinate-document-candidates-v1.json
  map/data/source-archive/concessions/nkr-attachments/<sha256>.<ext>
  map/data/source-archive/concessions/nkr-attachments-text/<sha256>.txt

Safety:
- downloading/extracting a document never publishes geometry;
- numeric tables are only classified as coordinate candidates;
- exact geometry still passes the baseline publication gate separately.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import mimetypes
import pathlib
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from urllib.parse import urljoin

import requests
import xlrd
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = pathlib.Path(__file__).resolve().parents[1]
DATA = ROOT / "map" / "data"
PARTIES_PATH = DATA / "nkr-party-documents-v1.json"
PENDING_PATH = DATA / "pending-concessions-v1.json"
CURATED_PATH = DATA / "concession-boundaries-official-source-v2.json"
INDEX_PATH = DATA / "nkr-attachment-index-v1.json"
CANDIDATES_PATH = DATA / "coordinate-document-candidates-v1.json"
RAW_DIR = DATA / "source-archive" / "concessions" / "nkr-attachments"
TEXT_DIR = DATA / "source-archive" / "concessions" / "nkr-attachments-text"

BASE = "https://nkr.government.bg"
LANDING = "/Concessions"
UA = "EnergoKarta-Bulgaria-NKR-attachments/1.0 (+https://topperbg.github.io/map/)"
MAX_NEW_FILES = 60
MAX_TOTAL_BYTES = 60 * 1024 * 1024
MAX_FILE_BYTES = 15 * 1024 * 1024
SLEEP = 0.8

PRIORITY_IDS = {
    "D-000074": 0,  # Elatsite
    "D-000034": 0,  # Chelopech
    "D-000327": 0, "D-000317": 0, "D-000324": 0,  # Bobov Dol
}

KEYWORDS = re.compile(
    r"координат|координати|граничн[аи]|регистър\s+на\s+координат|"
    r"бгс\s*2005|кс\s*1970|координатна\s+система|"
    r"concession\s+boundary|coordinate",
    re.I,
)

NUM_RE = re.compile(r"(?<![\w])[-+]?\d{1,2}(?:[ .\u00a0]\d{3})*(?:[.,]\d+)?|(?<![\w])[-+]?\d{3,10}(?:[.,]\d+)?")


def now_iso() -> str:
    return dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def load_json(path: pathlib.Path, default: dict) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def safe_number(token: str) -> float | None:
    s = token.replace("\u00a0", "").replace(" ", "")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def classify_pair(a: float, b: float) -> str | None:
    aa, bb = abs(a), abs(b)
    # BGS 1970 zones: X/N about 4.4-4.9M; Y/E about 8.3-9.7M.
    if (4.3e6 <= aa <= 5.0e6 and 8.0e6 <= bb <= 10.0e6) or (
        4.3e6 <= bb <= 5.0e6 and 8.0e6 <= aa <= 10.0e6
    ):
        return "BGS1970"
    # BGS2005 / CCS2005: northing in the 4.xM range and easting in the
    # few-hundred-km range. Keep deliberately broad; publication QA is stricter.
    if (4.3e6 <= aa <= 5.0e6 and 1.0e5 <= bb <= 1.2e6) or (
        4.3e6 <= bb <= 5.0e6 and 1.0e5 <= aa <= 1.2e6
    ):
        return "BGS2005"
    if (41.0 <= aa <= 44.5 and 22.0 <= bb <= 29.5) or (
        41.0 <= bb <= 44.5 and 22.0 <= aa <= 29.5
    ):
        return "WGS84"
    return None


def coordinate_rows(text: str) -> list[dict]:
    rows = []
    seen = set()
    for line_no, raw in enumerate(text.splitlines(), 1):
        line = re.sub(r"\s+", " ", raw).strip()
        if not line:
            continue
        nums = [safe_number(x) for x in NUM_RE.findall(line)]
        nums = [x for x in nums if x is not None]
        candidates = []
        # tables commonly look like: pointNo X Y, or X Y.
        if len(nums) >= 3:
            candidates.append((nums[-2], nums[-1]))
        if len(nums) >= 2:
            candidates.append((nums[0], nums[1]))
            candidates.append((nums[-2], nums[-1]))
        for a, b in candidates:
            crs = classify_pair(a, b)
            if not crs:
                continue
            key = (round(a, 4), round(b, 4), line_no)
            if key in seen:
                continue
            seen.add(key)
            rows.append({
                "line": line_no,
                "a": a,
                "b": b,
                "crsClass": crs,
                "text": line[:280],
            })
            break
    return rows


def infer_extension(headers: dict, data: bytes) -> tuple[str, str | None]:
    cd = headers.get("Content-Disposition", "")
    filename = None
    m = re.search(r"filename\*?=(?:UTF-8''|\")?([^\";]+)", cd, re.I)
    if m:
        filename = m.group(1).strip().strip('"')
    ext = pathlib.Path(filename or "").suffix.lower()
    ctype = (headers.get("Content-Type") or "").split(";")[0].strip().lower()
    if not ext:
        if data.startswith(b"%PDF"):
            ext = ".pdf"
        elif data.startswith(b"PK\x03\x04"):
            # distinguish OOXML by zip members
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    names = set(z.namelist())
                if "word/document.xml" in names:
                    ext = ".docx"
                elif "xl/workbook.xml" in names:
                    ext = ".xlsx"
                else:
                    ext = ".zip"
            except Exception:
                ext = ".zip"
        elif data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            ext = ".xls"
        elif "html" in ctype:
            ext = ".html"
        elif ctype.startswith("text/"):
            ext = ".txt"
        else:
            ext = mimetypes.guess_extension(ctype) or ".bin"
    return ext, filename


def extract_text(data: bytes, ext: str, *, allow_ocr: bool = False) -> tuple[str, str]:
    ext = ext.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            text = "\n\f\n".join((page.extract_text() or "") for page in reader.pages)
            method = "pypdf"
            if len(text.strip()) < 500 and shutil.which("pdftotext"):
                with tempfile.TemporaryDirectory() as td:
                    src = pathlib.Path(td) / "source.pdf"
                    out = pathlib.Path(td) / "source.txt"
                    src.write_bytes(data)
                    subprocess.run(["pdftotext", "-layout", str(src), str(out)], check=False, timeout=90)
                    alt = out.read_text(encoding="utf-8", errors="replace") if out.exists() else ""
                    if len(alt.strip()) > len(text.strip()):
                        text, method = alt, "pdftotext"
            if allow_ocr and len(text.strip()) < 500 and shutil.which("pdftoppm") and shutil.which("tesseract"):
                with tempfile.TemporaryDirectory() as td:
                    src = pathlib.Path(td) / "source.pdf"
                    src.write_bytes(data)
                    prefix = pathlib.Path(td) / "page"
                    subprocess.run(
                        ["pdftoppm", "-jpeg", "-r", "140", "-f", "1", "-l", "30", str(src), str(prefix)],
                        check=False, timeout=180, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
                    )
                    parts = []
                    for image in sorted(pathlib.Path(td).glob("page-*.jpg")):
                        try:
                            p = subprocess.run(
                                ["tesseract", str(image), "stdout", "-l", "bul+eng", "--psm", "6"],
                                check=False, timeout=60, capture_output=True
                            )
                            parts.append(p.stdout.decode("utf-8", errors="replace"))
                        except Exception:
                            continue
                    ocr = "\n\f\n".join(parts)
                    if len(ocr.strip()) > len(text.strip()):
                        text, method = ocr, "tesseract-bul+eng"
            return text, method
        if ext == ".doc":
            if shutil.which("antiword"):
                with tempfile.TemporaryDirectory() as td:
                    src = pathlib.Path(td) / "source.doc"
                    src.write_bytes(data)
                    p = subprocess.run(["antiword", str(src)], check=False, timeout=90, capture_output=True)
                    text = p.stdout.decode("utf-8", errors="replace")
                    if text.strip():
                        return text, "antiword"
            return "", "unsupported-doc"
        if ext == ".docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
            return "\n".join(parts), "python-docx"
        if ext == ".xlsx":
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"### SHEET {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v not in (None, "") for v in row):
                        lines.append("\t".join("" if v is None else str(v) for v in row))
            return "\n".join(lines), "openpyxl"
        if ext == ".xls":
            book = xlrd.open_workbook(file_contents=data)
            lines = []
            for sheet in book.sheets():
                lines.append(f"### SHEET {sheet.name}")
                for r in range(sheet.nrows):
                    vals = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
                    if any(v not in (None, "") for v in vals):
                        lines.append("\t".join(str(v) for v in vals))
            return "\n".join(lines), "xlrd"
        if ext in (".html", ".htm"):
            soup = BeautifulSoup(data, "lxml")
            return soup.get_text("\n"), "beautifulsoup"
        if ext in (".txt", ".csv", ".tsv", ".rtf"):
            for enc in ("utf-8", "cp1251", "latin-1"):
                try:
                    return data.decode(enc), f"decode-{enc}"
                except UnicodeDecodeError:
                    pass
    except Exception as exc:
        return "", f"extract-error:{type(exc).__name__}:{exc}"
    return "", "unsupported"


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept-Language": "bg-BG,bg;q=0.9,en;q=0.7"})
    r = s.get(BASE + LANDING, timeout=60)
    r.raise_for_status()
    time.sleep(SLEEP)
    return s


def expected_by_concession(curated: dict) -> dict[str, dict]:
    out = {}
    for rec in curated.get("records") or []:
        cid = rec.get("concession_registry")
        if not cid:
            continue
        expected = rec.get("point_count")
        if expected is None and rec.get("deposit_point_count") and rec.get("additional_area_point_count"):
            expected = rec["deposit_point_count"] + rec["additional_area_point_count"]
        out[cid] = {
            "id": rec.get("id"),
            "name": rec.get("name"),
            "expectedPointCount": expected,
            "sourceCrs": rec.get("source_coordinate_system"),
            "officialAreaDka": rec.get("official_area_dka"),
            "status": rec.get("status"),
        }
    return out


def main() -> None:
    parties_payload = load_json(PARTIES_PATH, {"parties": {}})
    pending_payload = load_json(PENDING_PATH, {"pending": []})
    curated = load_json(CURATED_PATH, {"records": []})
    old_index = load_json(INDEX_PATH, {"files": {}})
    files = dict(old_index.get("files") or {})

    expected = expected_by_concession(curated)
    priority_by_inventory = {}
    for item in pending_payload.get("pending") or []:
        iid = item.get("inventoryId") or item.get("concessionId")
        if iid:
            priority_by_inventory[iid] = min(
                priority_by_inventory.get(iid, 99), int(item.get("priority") or 4)
            )

    links = {}
    for party in (parties_payload.get("parties") or {}).values():
        inv = party.get("inventoryIds") or party.get("concessionIds") or []
        cids = party.get("concessionIds") or []
        base_priority = min([PRIORITY_IDS.get(cid, 9) for cid in cids] + [9])
        inv_priority = min([priority_by_inventory.get(i, 4) for i in inv] + [4])
        for link in party.get("fileLinks") or []:
            url = link.get("url")
            if not url:
                continue
            fid = url.rstrip("/").split("/")[-1]
            rec = links.setdefault(fid, {
                "fileId": fid,
                "url": url,
                "title": link.get("title"),
                "partyGuids": [],
                "inventoryIds": [],
                "concessionIds": [],
                "priority": min(base_priority, inv_priority),
            })
            rec["priority"] = min(rec["priority"], base_priority, inv_priority)
            if party.get("guid") not in rec["partyGuids"]:
                rec["partyGuids"].append(party.get("guid"))
            for x in inv:
                if x and x not in rec["inventoryIds"]:
                    rec["inventoryIds"].append(x)
            for x in cids:
                if x and x not in rec["concessionIds"]:
                    rec["concessionIds"].append(x)

    def needs_processing(link: dict) -> bool:
        old = files.get(link["fileId"])
        if not old or old.get("status") != "ok":
            return True
        extractor = str(old.get("extractor") or "")
        if extractor.startswith("unsupported") or extractor.startswith("extract-error"):
            return True
        if old.get("extension") == ".pdf" and int(old.get("textChars") or 0) < 500:
            return any(cid in PRIORITY_IDS for cid in link.get("concessionIds") or [])

        return False

    todo = [x for x in links.values() if needs_processing(x)]
    todo.sort(key=lambda x: (x["priority"], x["concessionIds"][0] if x["concessionIds"] else "", x["fileId"]))

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    TEXT_DIR.mkdir(parents=True, exist_ok=True)
    errors = []
    downloaded = 0
    downloaded_bytes = 0

    try:
        session = make_session()
    except Exception as exc:
        session = None
        errors.append({"stage": "session", "error": str(exc)})

    for item in todo:
        if downloaded >= MAX_NEW_FILES or downloaded_bytes >= MAX_TOTAL_BYTES or session is None:
            break
        try:
            r = session.get(item["url"], timeout=120)
            r.raise_for_status()
            data = r.content
            if not data:
                raise RuntimeError("empty attachment")
            if len(data) > MAX_FILE_BYTES:
                files[item["fileId"]] = {**item, "status": "skipped-too-large", "bytes": len(data), "checkedAt": now_iso()}
                continue
            ext, filename = infer_extension(r.headers, data)
            digest = hashlib.sha256(data).hexdigest()
            raw_path = RAW_DIR / f"{digest}{ext}"
            if not raw_path.exists():
                raw_path.write_bytes(data)
            text, extractor = extract_text(
                data, ext,
                allow_ocr=any(cid in PRIORITY_IDS for cid in item.get("concessionIds") or [])
            )
            text_path = TEXT_DIR / f"{digest}.txt"
            if text and not text_path.exists():
                text_path.write_text(text, encoding="utf-8")
            rows = coordinate_rows(text)
            counts = {}
            for row in rows:
                counts[row["crsClass"]] = counts.get(row["crsClass"], 0) + 1
            files[item["fileId"]] = {
                **item,
                "status": "ok",
                "checkedAt": now_iso(),
                "filename": filename,
                "extension": ext,
                "contentType": r.headers.get("Content-Type"),
                "bytes": len(data),
                "sha256": digest,
                "archivePath": str(raw_path.relative_to(DATA)).replace("\\", "/"),
                "textPath": str(text_path.relative_to(DATA)).replace("\\", "/") if text else None,
                "extractor": extractor,
                "textChars": len(text),
                "keywordHit": bool(KEYWORDS.search(text)),
                "coordinateRowCount": len(rows),
                "coordinateClassCounts": counts,
                "coordinateSample": rows[:12],
            }
            downloaded += 1
            downloaded_bytes += len(data)
        except Exception as exc:
            errors.append({"fileId": item["fileId"], "url": item["url"], "error": str(exc)})
            previous = files.get(item["fileId"]) or {}
            files[item["fileId"]] = {**previous, **item, "status": previous.get("status") or "error", "lastError": str(exc), "lastErrorAt": now_iso()}
        time.sleep(SLEEP)

    candidates = []
    for rec in files.values():
        if rec.get("status") != "ok":
            continue
        row_count = int(rec.get("coordinateRowCount") or 0)
        if not rec.get("keywordHit") and row_count < 4:
            continue
        exp = None
        for cid in rec.get("concessionIds") or []:
            if cid in expected:
                exp = {"concessionId": cid, **expected[cid]}
                break
        counts = rec.get("coordinateClassCounts") or {}
        dominant = max(counts, key=counts.get) if counts else None
        exact_expected = bool(exp and exp.get("expectedPointCount") and counts.get(dominant) == exp["expectedPointCount"])
        candidates.append({
            "fileId": rec.get("fileId"),
            "sha256": rec.get("sha256"),
            "archivePath": rec.get("archivePath"),
            "textPath": rec.get("textPath"),
            "inventoryIds": rec.get("inventoryIds") or [],
            "concessionIds": rec.get("concessionIds") or [],
            "filename": rec.get("filename"),
            "extension": rec.get("extension"),
            "textChars": rec.get("textChars"),
            "keywordHit": rec.get("keywordHit"),
            "coordinateRowCount": row_count,
            "coordinateClassCounts": counts,
            "dominantCrsClass": dominant,
            "expected": exp,
            "exactExpectedPointCount": exact_expected,
            "qaStatus": "ready-for-coordinate-parser" if exact_expected else "candidate-review",
            "sample": rec.get("coordinateSample") or [],
        })

    candidates.sort(key=lambda x: (
        0 if x.get("exactExpectedPointCount") else 1,
        -(x.get("coordinateRowCount") or 0),
        (x.get("concessionIds") or [""])[0],
    ))

    index_payload = {
        "schema": "bgwf-nkr-attachment-index-v1",
        "generatedAt": now_iso(),
        "runtimeDependency": False,
        "knownLinkCount": len(links),
        "archivedOkCount": sum(1 for x in files.values() if x.get("status") == "ok"),
        "downloadedThisRun": downloaded,
        "downloadedBytesThisRun": downloaded_bytes,
        "errors": errors,
        "files": files,
    }
    INDEX_PATH.write_text(json.dumps(index_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    candidate_payload = {
        "schema": "bgwf-coordinate-document-candidates-v1",
        "generatedAt": now_iso(),
        "candidateCount": len(candidates),
        "readyForCoordinateParserCount": sum(1 for x in candidates if x.get("qaStatus") == "ready-for-coordinate-parser"),
        "rules": [
            "Attachment extraction alone never publishes geometry.",
            "A ready-for-coordinate-parser candidate matches a curated expected point count in one detected CRS class.",
            "Publication still requires CRS/order verification, Bulgaria bounds and official-area QA.",
        ],
        "candidates": candidates,
    }
    CANDIDATES_PATH.write_text(json.dumps(candidate_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"NKR attachments: links={len(links)}; archived={index_payload['archivedOkCount']}; "
        f"downloaded={downloaded} ({downloaded_bytes} bytes); candidates={len(candidates)}; "
        f"ready={candidate_payload['readyForCoordinateParserCount']}; errors={len(errors)}"
    )


if __name__ == "__main__":
    main()
